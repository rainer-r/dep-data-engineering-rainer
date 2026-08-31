"""Validator for the BIR tax-collection source (Excel workbooks).

Contract: ``data/reference/schemas/bir_tax.json``. The workbook layout is
self-describing: a title row, a period row, a units row, then a header row
whose first cell is ``PARTICULARS`` followed by period columns and a
``TOTAL`` column. Registered files come from the committed inventory in the
schema artifact; the manifest check lives in ``global_checks``.
"""

from __future__ import annotations

import openpyxl

from src.validation.base import (
    SEVERITY_MUST,
    SEVERITY_SHOULD,
    BaseValidator,
    ValidationReport,
    ValidationResult,
)

_MONTH_TOKENS = [
    "JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE",
    "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER",
]


def _norm(value) -> str:
    """Normalise a cell value to its string form for token comparison."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_header_row(ws) -> list:
    """Return the tokenised row whose first cell is 'PARTICULARS'."""
    for row in ws.iter_rows(values_only=True):
        vals = [_norm(v) for v in row]
        if vals and vals[0] == "PARTICULARS":
            return vals
    return []


class TaxCollectionValidator(BaseValidator):
    """Validates data/raw/tax/ against the BIR expected-schema artifact."""

    source = "tax"
    schema_name = "bir_tax"
    source_label = "BIR Tax Collection (PSCG)"

    # -- layout helpers -----------------------------------------------------

    def _sheet_layout(self, ws, title_tokens, period_tokens, units_tokens,
                      must_tokens, min_months=0) -> str:
        """Return '' when the sheet satisfies the layout contract else a reason."""
        header = _find_header_row(ws)
        if not header:
            return "no PARTICULARS header row found"

        rows = list(ws.iter_rows(values_only=True))
        text = " ".join(_norm(v) for row in rows for v in row).upper()
        for label, tokens in (("title", title_tokens), ("period", period_tokens),
                              ("units", units_tokens)):
            missing = [t for t in tokens if t.upper() not in text]
            if missing:
                return "%s token(s) missing: %s" % (label, missing)
        if "TOTAL" in must_tokens and "TOTAL" not in header:
            return "header row lacks a TOTAL token"
        for token in must_tokens - {"TOTAL"}:
            if token not in header:
                return "header row lacks token %s" % token
        found_months = [m for m in _MONTH_TOKENS if m in header]
        if len(found_months) < min_months:
            return "header row has only %d month token(s)" % len(found_months)
        return ""

    def _check_file(self, rep, rel_path: str, spec: dict):
        path = self.raw_dir / rel_path
        if not path.exists():
            rep.add(ValidationResult("tax.registered_files_present", SEVERITY_MUST, False,
                                     detail="missing registered workbook %s" % rel_path,
                                     actual=rel_path, expected="on disk"))
            return

        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 — captures corrupt/truncated workbooks
            rep.add(ValidationResult("tax.registered_openable", SEVERITY_MUST, False,
                                     detail="cannot open %s: %s" % (rel_path, e),
                                     actual=rel_path, expected="openable xlsx"))
            return

        sheet_names = set(wb.sheetnames)
        expected_sheets = {str(s) for s in spec["sheets"]}
        sheets_ok = sheet_names == expected_sheets
        rep.add(
            ValidationResult(
                "tax.sheet_expectations", SEVERITY_MUST, sheets_ok,
                detail="%s sheets %s vs expected %s" % (rel_path, sorted(sheet_names), sorted(expected_sheets)),
                actual=sorted(sheet_names), expected=sorted(expected_sheets),
            )
        )

        layout = spec.get("layout", {})
        title = layout.get("title", ["INTERNAL REVENUE COLLECTIONS"])
        period = layout.get("period", ["For the Period:"])

        first_data_sheet = None
        for name in spec["sheets"]:
            if str(name) in wb.sheetnames:
                first_data_sheet = wb[str(name)]
                break
        if first_data_sheet is None:
            wb.close()
            return

        kind = spec["kind"]
        min_months = 2 if kind in ("monthly", "partial-year") else 0
        must_tokens = {"TOTAL"} if kind in ("monthly", "partial-year") else set()
        lay = self._sheet_layout(first_data_sheet, title, period, ["Million Pesos"], must_tokens,
                                 min_months=min_months)
        rep.add(
            ValidationResult(
                "tax.layout_contract", SEVERITY_MUST, not lay,
                detail="%s: %s" % (rel_path, lay or "layout matches contract"),
                actual=rel_path, expected="title/period/units/PARTICULARS/TOTAL present",
            )
        )
        self._kind_specific_checks(rep, wb, kind)
        wb.close()

    # -- kind-specific checks ----------------------------------------------

    def _kind_specific_checks(self, rep, wb, kind: str):
        if kind == "annual":
            year_sheet = wb["2005_2024"] if "2005_2024" in wb.sheetnames else None
            if year_sheet is not None:
                header = _find_header_row(year_sheet)
                has_year = any(
                    t.isdigit() and 1990 <= int(t) <= 2100 for t in header
                )
                rep.add(
                    ValidationResult(
                        "tax.annual_year_columns", SEVERITY_SHOULD, has_year,
                        detail="annual sheet header includes calendar-year columns (e.g. 2005..2024)",
                        actual=has_year, expected=True,
                    )
                )
            total_any = any(
                (_find_header_row(ws) and "TOTAL" in _find_header_row(ws))
                for ws in wb.worksheets
            )
            rep.add(
                ValidationResult(
                    "tax.total_column_present", SEVERITY_SHOULD, total_any,
                    detail="at least one sheet exposes a TOTAL column",
                    actual=total_any, expected=True,
                )
            )
        elif kind == "partial-year":
            notes = wb["Notes"] if "Notes" in wb.sheetnames else None
            if notes is not None:
                notes_txt = " ".join(
                    _norm(v) for row in notes.iter_rows(values_only=True) for v in row
                )
                has_source = "Revenue Accounting Division" in notes_txt
                rep.add(
                    ValidationResult(
                        "tax.notes_sheet_attribution", SEVERITY_SHOULD, has_source,
                        detail="Notes sheet attributes source to Revenue Accounting Division, BIR",
                        actual=has_source, expected=True,
                    )
                )

    # -- main ---------------------------------------------------------------

    def validate(self) -> ValidationReport:
        schema = self.load_schema()
        rep = self.report()
        rep.add(self.register_schema_check())

        registered: list = schema["files"]

        # registered files present -----------------------------------------
        missing = [spec["relative_path"] for spec in registered
                   if not (self.raw_dir / spec["relative_path"]).exists()]
        rep.add(
            ValidationResult(
                "tax.registered_files_present", SEVERITY_MUST, not missing,
                detail="missing registered workbooks: %s" % (missing or "none"),
                actual=len(registered) - len(missing), expected=len(registered),
            )
        )

        # per-file layout + sheet checks ------------------------------------
        for spec in registered:
            self._check_file(rep, spec["relative_path"], spec)

        # unregistered files on disk (completeness) -------------------------
        registered_keys = {spec["relative_path"] for spec in registered}
        tax_dir = self.raw_dir / "tax"
        unregistered = sorted(
            "tax/%s" % p.name
            for p in sorted(tax_dir.glob("*.xlsx"))
            if p.is_file() and "tax/%s" % p.name not in registered_keys
        )
        rep.add(
            ValidationResult(
                "tax.unregistered_files", SEVERITY_SHOULD, not unregistered,
                detail=("workbook(s) present on disk but not in the committed inventory: %s"
                        % (unregistered or "none")),
                actual=unregistered or None, expected="[]",
            )
        )
        return rep
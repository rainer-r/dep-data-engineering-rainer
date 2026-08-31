"""Shared fixtures for the validation-layer test suite.

Fixtures build small synthetic raw-data trees in temp dirs (fully offline)
with the *real* committed expected-schema artifacts as the contract. The
anti-tautology rule: expected literals in assertions are written out as
constants, never read back from the artifact being validated.
"""

from __future__ import annotations

from pathlib import Path

import hashlib
import json

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
import shutil

from src.validation.ph_gaa import GAAValidator
from src.validation.ph_saaodb import SAAODBValidator
from src.validation.ph_tax_collection import TaxCollectionValidator
from src.validation.global_checks import GlobalValidator

REPO_ROOT = Path(__file__).resolve().parents[1]
SCHEMAS_DIR = REPO_ROOT / "data/reference/schemas"

GAA_COLUMNS = [
    "id", "department", "uacs_dpt_dsc", "agency", "uacs_agy_dsc",
    "prexc_fpap_id", "dsc", "operunit", "uacs_oper_dsc", "uacs_reg_id",
    "fundcd", "uacs_fundsubcat_dsc", "uacs_exp_cd", "uacs_exp_dsc",
    "uacs_sobj_cd", "uacs_sobj_dsc", "amt", "year", "prexc_level",
    "sorder", "uacs_operdiv_id", "uacs_div_dsc",
]
GAA_NULLABLE = {"prexc_level", "sorder", "uacs_operdiv_id", "uacs_div_dsc"}

MONTHS = ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE",
          "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"]


# ---- GAA mini parquet -----------------------------------------------------

def _mini_gaa_table(n: int, years, exp_codes, reg_ids, neg_count: int) -> pa.Table:
    data = {}
    for name in GAA_COLUMNS:
        if name == "id":
            vals = list(range(1, n + 1))
        elif name == "year":
            vals = [years[i % len(years)] for i in range(n)]
        elif name == "amt":
            vals = [1000.0 + float(i) for i in range(n)]
            for i in range(min(neg_count, n)):
                vals[i] = -vals[i]
        elif name == "uacs_exp_cd":
            vals = [exp_codes[i % len(exp_codes)] for i in range(n)]
        elif name == "uacs_reg_id":
            vals = [reg_ids[i % len(reg_ids)] for i in range(n)]
        elif name in GAA_NULLABLE:
            vals = [None for _ in range(n)]
        else:
            vals = [f"{name}-0{i}" for i in range(n)]
        if name in ("id", "year"):
            data[name] = pa.array(vals, type=pa.int64())
        elif name == "amt":
            data[name] = pa.array(vals, type=pa.float64())
        else:
            data[name] = pa.array(vals, type=pa.string())
    return pa.table(data)


def make_gaa_tree(root: Path, n: int = 3, *, years=None, exp_codes=None,
                  reg_ids=None, neg_count: int = 0) -> Path:
    """Build gaa/gaa.parquet + gaa/README.md under root; return root."""
    gaa_dir = root / "gaa"
    gaa_dir.mkdir(parents=True, exist_ok=True)
    pq.write_table(
        _mini_gaa_table(
            n,
            years or [2020, 2021, 2022, 2023, 2024, 2025, 2026],
            exp_codes or ["1", "2", "3", "6", "nan"],
            reg_ids or ["00", "nan", "15", "18", "01"],
            neg_count,
        ),
        gaa_dir / "gaa.parquet",
    )
    (gaa_dir / "README.md").write_text(
        "# GAA dataset\n\nDownloaded from bettergovph/gaa on Hugging Face (2026-08-30).",
        encoding="utf-8",
    )
    return root

# ---- tiny BIR xlsx --------------------------------------------------------

def _write_contract_rows(ws, header, *, title="INTERNAL REVENUE COLLECTIONS",
                         period="For the Period: CY 2025", units="(Amount in Million Pesos)",
                         body=None):
    ws.append([f"{title} BY REGION AND PROVINCE"])
    ws.append([period])
    ws.append([units])
    ws.append(header)
    for row in body or []:
        ws.append(row)


def _good_annual(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2005_2024"
    _write_contract_rows(
        ws, ["PARTICULARS"] + [str(y) for y in range(2005, 2025)] + ["TOTAL"],
        period="For the Period: CY 2005 - CY 2024",
        body=[["National Capital Region"] + [1000] * 20 + [20000]])
    ws2 = wb.create_sheet("2025")
    _write_contract_rows(
        ws2, ["PARTICULARS"] + MONTHS + ["TOTAL"],
        period="For the Period: CY 2025",
        body=[["National Capital Region"] + [100] * 12 + [1200]])
    wb.save(path)


def _good_monthly(path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for y in range(2016, 2026):
        ws = wb.create_sheet(str(y))
        _write_contract_rows(
            ws, ["PARTICULARS"] + MONTHS + ["TOTAL"],
            period=f"For the Period: CY {y}",
            body=[["National Capital Region"] + [10] * 12 + [120]])
    wb.save(path)


def _good_partial_year(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PSA-PSGC"
    _write_contract_rows(
        ws, ["PARTICULARS"] + MONTHS + ["TOTAL"],
        period="For the Period: January - May 2026",
        body=[["National Capital Region"] + [10] * 12 + [120]])
    notes = wb.create_sheet("Notes")
    notes.append(["Source: Revenue Accounting Division, BIR"])
    wb.save(path)


def _broken_partial_year(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PSA-PSGC"
    _write_contract_rows(
        ws, ["PARTICULARS", "2026"],
        period="For the Period: 2026",
        body=[["National Capital Region", 1]])
    notes = wb.create_sheet("Notes")
    notes.append(["Something else"])
    wb.save(path)


TAX_FILENAMES = {
    "annual": "20260430_Annual_Collection_PSCG_2005_2025 (website) cv.xlsx",
    "monthly": "20260430_Monthly_Collection_PSCG_2016_2025 (website) cv.xlsx",
    "partial": "20260616_Monthly_Collection_PSCG_JanMay2026 (D).xlsx",
}


def make_tax_tree(root: Path) -> Path:
    tax = root / "tax"
    tax.mkdir(parents=True, exist_ok=True)
    _good_annual(tax / TAX_FILENAMES["annual"])
    _good_monthly(tax / TAX_FILENAMES["monthly"])
    _good_partial_year(tax / TAX_FILENAMES["partial"])
    return root


# ---- pytest fixtures ------------------------------------------------------

@pytest.fixture()
def gaa_raw_dir(tmp_path):
    return make_gaa_tree(tmp_path)


@pytest.fixture()
def gaa_validator(gaa_raw_dir):
    return GAAValidator(raw_dir=gaa_raw_dir, schema_dir=SCHEMAS_DIR)


@pytest.fixture()
def tax_raw_dir(tmp_path):
    return make_tax_tree(tmp_path)


@pytest.fixture()
def tax_validator(tax_raw_dir):
    return TaxCollectionValidator(raw_dir=tax_raw_dir, schema_dir=SCHEMAS_DIR)


@pytest.fixture()
def by_id():
    """Return a mapper: report -> {check_id: [ValidationResult]} (ids repeat)."""
    def mapper(rep):
        out = {}
        for r in rep.results:
            out.setdefault(r.check, []).append(r)
        return out
    return mapper


# ---- SAAODB fixture helpers ----------------------------------------------

def make_one_page_pdf(path: Path) -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with open(path, "wb") as f:
        writer.write(f)


def make_saaodb_tree(root: Path, *, good_htm: bool = True, with_pdf: bool = True) -> Path:
    """Build saaodb/ with one .htm stub (committed fixture) + optional 1-page PDF."""
    d = root / "saaodb"
    d.mkdir(parents=True, exist_ok=True)
    stub = "saaodb_good.htm" if good_htm else "saaodb_bad.htm"
    shutil.copy(REPO_ROOT / "tests/fixtures" / stub, d / "2011Q1.htm")
    if with_pdf:
        make_one_page_pdf(d / "2026Q1.pdf")
    return root


@pytest.fixture()
def saaodb_raw_dir(tmp_path):
    return make_saaodb_tree(tmp_path)


@pytest.fixture()
def saaodb_validator(saaodb_raw_dir):
    return SAAODBValidator(raw_dir=saaodb_raw_dir, schema_dir=SCHEMAS_DIR,
                           skip_live=True)

# ---- global validator (shared manifest + raw-data tree) --------------------

GLOBAL_FILES = {
    "gaa/gaa.parquet": b"PAR1" + b"\x00" * 128,
    "gaa/README.md": b"# GAA fixture dataset",
    "tax/20260430_Annual_Collection_PSCG_2005_2025 (website) cv.xlsx":
        b"PK\x03\x04" + b"\x00" * 64,
    "saaodb/2026_Q1_1st_Quarter.htm":
        b"<html><body><table><tr><td>ALLOTMENT</td></tr></table></body></html>",
    "saaodb/2026_Q2_2nd_Quarter.pdf": b"%PDF-1.4\n%%EOF",
}
GLOBAL_UNREGISTERED = [
    "tax/20260716_Monthly_Collection_PSCG_JanJun2026 (D) upload.xlsx",
]


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def make_manifest_tree(root: Path, files: dict, *, extra=None,
                       corrupt_rel=None) -> Path:
    """Write registered files - and optional unregistered 'extra' ones - under
    root, then a manifest.json whose entries carry each file's real SHA-256.

    corrupt_rel: after writing, flip one byte of that registered file so its
    on-disk checksum no longer matches its manifest entry.
    """
    entries = {}
    for rel, data in files.items():
        p = root / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(data)
        entries[rel] = {
            "relative_path": rel,
            "checksum": _sha256_bytes(data),
            "size_bytes": len(data),
        }
    for rel in extra or []:
        p = root / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(b"unregistered fixture content")
    if corrupt_rel:
        p = root / corrupt_rel
        body = bytearray(p.read_bytes())
        body[0] ^= 0xFF
        p.write_bytes(bytes(body))
    (root / "manifest.json").write_text(
        json.dumps(entries, indent=2), encoding="utf-8")
    return root


@pytest.fixture()
def global_raw_dir(tmp_path):
    root = make_manifest_tree(tmp_path, GLOBAL_FILES, extra=GLOBAL_UNREGISTERED)
    # hidden entries that must never surface in the unregistered scan
    (root / "gaa/.cache").mkdir(parents=True, exist_ok=True)
    (root / "gaa/.cache/stale.parquet").write_bytes(b"stale")
    (root / "gaa/.gitkeep").write_bytes(b"")
    return root


@pytest.fixture()
def global_validator(global_raw_dir):
    return GlobalValidator(raw_dir=global_raw_dir,
                           manifest_file=global_raw_dir / "manifest.json")
